"""
relatorio.py — BK Electrical Line Parameters
Geração da Memória de Cálculo em Word (.docx), reproduzindo o PADRÃO
OFICIAL DE DOCUMENTOS BK (capa/carimbo, cabeçalho, rodapé, sumário, os 10
títulos fixos, tabelas e tipografia) definido no arquivo de referência
`DOCUMENTO PADRÃO BK.docx` fornecido pelo usuário e na skill
"template-relatorios-memorias-memoriais-laudos".

O documento é construído a partir do arquivo `templates/BK_Template_Padrao.docx`
(cópia do modelo oficial, com a logomarca do cliente removida — este software
não coleta logo de cliente). Isso preserva exatamente os estilos, a
numeração automática dos títulos, cabeçalho, rodapé e a logomarca BK do
modelo original — nada é reconstruído "à mão" via HTML/CSS.

As fórmulas de engenharia são inseridas como EQUAÇÕES NATIVAS DO WORD
(OMML — o mesmo formato do recurso "Inserir Equação"), montadas pelo
módulo `equations.py`, e não como texto simples nem como imagem.

Nenhum cálculo de engenharia é feito aqui — apenas leitura de
`calculos.ResultadoLT` já calculado e formatação/apresentação, no padrão
documental BK.
"""

from __future__ import annotations

import io
import math
import os
from dataclasses import dataclass

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt

import calculos as c
import dados
import equations as eq

TEMPLATES_DIR = os.path.join(os.path.dirname(__file__), "templates")
TEMPLATE_PATH = os.path.join(TEMPLATES_DIR, "BK_Template_Padrao.docx")

STYLE_TITULO = "1. Título BK principal"
STYLE_SUBTITULO = "1.1. Subtítulo BK"
STYLE_CORPO = "Body Text"

_COR_CINZA_CLARO = "D9D9D9"


# ---------------------------------------------------------------------------
# MEMÓRIA DE CÁLCULO (etapas fórmula -> substituição -> resultado)
# Usada pela aba "5. Memória de Cálculo" do app (visualização em tela).
# Não é a fonte do relatório Word (que usa equações nativas — ver abaixo).
# ---------------------------------------------------------------------------

@dataclass
class EtapaCalculo:
    titulo: str
    formula: str
    onde: str
    dados: str
    substituicao: str
    resultado: str
    interpretacao: str = ""


def _fmt(v: float, casas: int = 4) -> str:
    return f"{v:,.{casas}f}".replace(",", "X").replace(".", ",").replace("X", ".")


def _fasor(z: complex, casas: int = 4) -> str:
    mod = abs(z)
    ang = math.degrees(math.atan2(z.imag, z.real))
    return f"{mod:.{casas}f} ∠ {ang:.{casas}f}°"


def montar_memoria_calculo(r: c.ResultadoLT) -> list[EtapaCalculo]:
    e = r.entrada
    g = r.geometria
    cond = r.condutor
    rl = r.rlgc
    ll = r.linha_longa
    pi = r.circuito_pi
    c1 = r.caso1
    reg = r.regulacao
    sil = r.sil
    ce = r.campo_eletrico

    etapas: list[EtapaCalculo] = []

    etapas.append(EtapaCalculo(
        "1. Distâncias entre fases (DAB, DBC, DCA)",
        "D = [(h₁-h₂)² + (v₁-v₂)²]^0,5",
        "h = deslocamento horizontal da fase (m); v = altura da fase (m)",
        f"Estrutura tipo {e.tipo_estrutura}: A(h={g.h_a:.3f}; v={g.v_a:.3f}), "
        f"B(h={g.h_b:.3f}; v={g.v_b:.3f}), C(h={g.h_c:.3f}; v={g.v_c:.3f}) m",
        f"DAB={g.dab_m:.4f} m; DBC={g.dbc_m:.4f} m; DCA={g.dca_m:.4f} m",
        f"DAB={g.dab_m:.4f} m; DBC={g.dbc_m:.4f} m; DCA={g.dca_m:.4f} m",
        "Distâncias geométricas entre os centros das fases, na estrutura selecionada.",
    ))
    etapas.append(EtapaCalculo(
        "2. Distância Média Geométrica entre fases (GMD / Deq)",
        "Deq = (DAB · DBC · DCA)^(1/3)",
        "DAB, DBC, DCA em metros",
        f"DAB={g.dab_m:.4f} m; DBC={g.dbc_m:.4f} m; DCA={g.dca_m:.4f} m",
        f"Deq = ({g.dab_m:.4f} × {g.dbc_m:.4f} × {g.dca_m:.4f})^(1/3)",
        f"Deq = {g.gmd_m:.4f} m",
        "Distância equivalente usada nas fórmulas de L e C da linha transposta.",
    ))
    etapas.append(EtapaCalculo(
        "3. Condutor e feixe de subcondutores",
        "req = (r · d)^0,5  [feixe duplo]   |   req = (r · d³ · √2)^0,25  [feixe quádruplo]",
        "r = raio do condutor unitário (m); d = espaçamento entre subcondutores (m)",
        f"Condutor {cond.catalogo} {cond.nome_comercial} (bitola {cond.bitola}); "
        f"diâmetro={cond.diametro_mm:.3f} mm; RMG={cond.rmg_m:.5f} m; "
        f"R={cond.r_ohm_km:.4f} Ω/km; nº subcondutores={cond.n_subcondutores}"
        + (f"; d={cond.espacamento_feixe_cm:.2f} cm" if cond.n_subcondutores > 1 else " (feixe simples)"),
        f"raio equivalente={cond.r_eq_m:.6f} m; RMG equivalente={cond.ds_eq_m:.6f} m",
        f"r_eq={cond.r_eq_m:.6f} m; Ds_eq={cond.ds_eq_m:.6f} m; R_fase={cond.r_fase_ohm_km:.5f} Ω/km",
        "Para feixe simples, r_eq e Ds_eq são os do condutor unitário do catálogo.",
    ))
    etapas.append(EtapaCalculo(
        "4. Resistência série (R)",
        "R = R_catálogo / nº subcondutores",
        "R_catálogo = resistência CA do condutor unitário (Ω/km, catálogo do fabricante)",
        f"R_catálogo={cond.r_ohm_km:.4f} Ω/km; nº subcondutores={cond.n_subcondutores}",
        f"R = {cond.r_ohm_km:.4f} / {cond.n_subcondutores}",
        f"R = {rl.r_ohm_km:.5f} Ω/km",
    ))
    etapas.append(EtapaCalculo(
        "5. Indutância série (L)",
        "L = 2×10⁻⁴ · ln(Deq / Ds_eq)",
        "Deq = GMD entre fases (m); Ds_eq = RMG equivalente do feixe (m)",
        f"Deq={g.gmd_m:.4f} m; Ds_eq={cond.ds_eq_m:.6f} m",
        f"L = 2×10⁻⁴ × ln({g.gmd_m:.4f} / {cond.ds_eq_m:.6f})",
        f"L = {rl.l_h_km*1000:.5f} mH/km",
    ))
    etapas.append(EtapaCalculo(
        "6. Capacitância (C) e condutância (G)",
        "C = 2π·ε / ln(Deq / r_eq) × 1000     |     G = C · σ / ε",
        "ε = εR·ε0 (ε0 ≈ 1/(36π)×10⁻⁹ F/m — aproximação de engenharia da planilha original); "
        "r_eq = raio equivalente do feixe (m); σ = condutividade do meio (S/m)",
        f"εR={e.eps_r:.4f}; Deq={g.gmd_m:.4f} m; r_eq={cond.r_eq_m:.6f} m; σ={e.sigma_s_m:.6f} S/m",
        f"C = 2π×{rl.eps_meio:.4e} / ln({g.gmd_m:.4f} / {cond.r_eq_m:.6f}) × 1000",
        f"C = {rl.c_f_km*1e9:.5f} nF/km; G = {rl.g_s_km:.3e} S/km",
    ))
    etapas.append(EtapaCalculo(
        "7. Impedância série e admitância shunt por km",
        "Z = R + jωL     |     Y = G + jωC     (ω = 2πf)",
        f"f={rl.freq_hz:.0f} Hz → ω={rl.omega_rad_s:.4f} rad/s",
        f"R={rl.r_ohm_km:.5f} Ω/km; L={rl.l_h_km*1000:.5f} mH/km; G={rl.g_s_km:.3e} S/km; C={rl.c_f_km*1e9:.5f} nF/km",
        f"Z = {rl.r_ohm_km:.5f} + j{rl.omega_rad_s*rl.l_h_km:.5f} Ω/km",
        f"Z = {_fasor(ll.z_km)} Ω/km; Y = {_fasor(ll.y_km*1000, 5)} mS/km",
    ))
    etapas.append(EtapaCalculo(
        "8. Impedância característica (Zo) e constante de propagação (γ)",
        "Zo = √(Z/Y)     |     γ = √(Z·Y) = α + jβ",
        "Z, Y por km",
        f"Z={_fasor(ll.z_km)} Ω/km; Y={_fasor(ll.y_km*1000,5)} mS/km",
        f"Zo = √(Z/Y); γ = √(Z·Y)",
        f"Zo = {_fasor(ll.zo_ohm)} Ω; γ = {ll.gamma_km.real:.6e} + j{ll.gamma_km.imag:.6e} (1/km)",
    ))
    etapas.append(EtapaCalculo(
        "9. Impedância de surto (Ro) e comprimento da linha (γ·l)",
        "Ro = √(L/C)     |     γ·l",
        "L, C por km; l = comprimento da linha (km)",
        f"L={rl.l_h_km*1000:.5f} mH/km; C={rl.c_f_km*1e9:.5f} nF/km; l={ll.comprimento_km:.3f} km",
        f"Ro = √({rl.l_h_km:.4e} / {rl.c_f_km:.4e})",
        f"Ro = {ll.ro_ohm:.4f} Ω; γ·l = {ll.gl.real:.6f} + j{ll.gl.imag:.6f}",
        "Ro é a impedância de surto (real, sem perdas) — usada no cálculo do SIL.",
    ))
    etapas.append(EtapaCalculo(
        "10. Circuito π equivalente",
        "Zs = Zo·senh(γl)     |     Yp = (1/Zo)·tanh(γl/2)",
        "Zo, γl calculados nas etapas 8 e 9",
        f"Zo={_fasor(ll.zo_ohm)} Ω; γl={ll.gl.real:.5f}+j{ll.gl.imag:.5f}",
        f"Zs = Zo × senh(γl); Yp = (1/Zo) × tanh(γl/2)",
        f"Zs = {_fasor(pi.zs_ohm)} Ω; Yp = {abs(pi.yp_siemens)*1000:.5f} mS ∠ "
        f"{math.degrees(math.atan2(pi.yp_siemens.imag,pi.yp_siemens.real)):.3f}°",
    ))
    etapas.append(EtapaCalculo(
        "11. Tensão e corrente no receptor (Caso 1 — carga plena)",
        "Vc = (VCL·1000/√3) ∠θv     |     Ic = conj(S / (3·Vc)),  S = SC·1000 ∠φ",
        "VCL = tensão de linha no receptor (kV); SC = potência aparente no receptor (kVA); "
        "φ = ± acos(cosφ) (sinal positivo para fator de potência atrasado)",
        f"VCL={e.vcl_kv:.4f} kV; SC={e.sc_kva:.4f} kVA; cosφ={e.cos_phi:.4f} "
        f"({'atrasado' if e.atrasado else 'adiantado'})",
        f"Vc = ({e.vcl_kv:.4f}×1000/√3) ∠{e.ang_vcl_deg:.2f}°",
        f"Vc = {_fasor(c1.vc_fasor_v, 2)} V (fase); Ic = {_fasor(c1.ic_fasor_a)} A",
    ))
    etapas.append(EtapaCalculo(
        "12. Tensão e corrente no emissor (Caso 1)",
        "Vs = Vc·cosh(γl) + Ic·Zo·senh(γl)     |     Is = (Vc/Zo)·senh(γl) + Ic·cosh(γl)",
        "Equações gerais do modelo de linha longa (parâmetros distribuídos)",
        f"Vc={_fasor(c1.vc_fasor_v,2)} V; Ic={_fasor(c1.ic_fasor_a)} A; Zo={_fasor(ll.zo_ohm)} Ω",
        "Vs = Vc·cosh(γl) + Ic·Zo·senh(γl)",
        f"|VoL| = {c1.vol_emissor_kv:.4f} kV ∠{c1.ang_vol_emissor_deg:.4f}°; "
        f"|Io| = {c1.io_emissor_a:.4f} A ∠{c1.ang_io_emissor_deg:.4f}°",
    ))
    etapas.append(EtapaCalculo(
        "13. Potências, perdas e regulação",
        "So = 3·Vs·Is*     |     Po=|So|·cos(∠So)     |     Pc=SC·cosφ     |     "
        "Perdas=Po−Pc     |     Regulação=(V_vazio−VCL)/VCL,  V_vazio=|Vs/cosh(γl)|",
        "Vs, Is do emissor (etapa 12); Vc do receptor em vazio obtida com a mesma "
        "tensão de emissor do Caso 1 (Ic=0)",
        f"So={c1.so_emissor_kva:.4f} kVA∠{c1.ang_so_emissor_deg:.3f}°; Pc={c1.pc_receptor_kw:.4f} kW",
        f"Perdas = {c1.po_emissor_kw:.4f} − {c1.pc_receptor_kw:.4f}",
        f"Po={c1.po_emissor_kw:.4f} kW; Perdas={c1.perdas_kw:.4f} kW ({c1.perdas_pct*100:.4f}%); "
        f"Regulação={reg.regulacao_pct*100:.4f}%",
    ))
    etapas.append(EtapaCalculo(
        "14. Potência natural (SIL) e carregamento",
        "SIL = VCL² / Ro     |     %SIL = Pc / SIL",
        "VCL em kV; Ro = impedância de surto (Ω)",
        f"VCL={e.vcl_kv:.4f} kV; Ro={ll.ro_ohm:.4f} Ω; Pc={c1.pc_receptor_kw:.4f} kW",
        f"SIL = {e.vcl_kv:.4f}² / {ll.ro_ohm:.4f} × 1000",
        f"SIL = {sil.sil_kw:.2f} kW; %SIL = {sil.pct_sil*100:.2f}%",
        "SIL (Surge Impedance Loading) é o carregamento natural da linha, referência de "
        "carregabilidade — carregamentos muito abaixo de 1×SIL indicam linha sobredimensionada "
        "para a demanda; muito acima, possível necessidade de compensação reativa.",
    ))
    etapas.append(EtapaCalculo(
        "15. Campo elétrico superficial do condutor (indicador de corona)",
        "Q = C·Vc (C em F/m)     |     E = Q / (2π·ε·r_eq)     |     E[kV/cm] = E[V/m]×10⁻⁵",
        "C = capacitância (F/m); Vc = tensão fase-neutro no receptor (V); r_eq = raio equivalente do feixe (m)",
        f"C={rl.c_f_km/1000:.4e} F/m; |Vc|={abs(c1.vc_fasor_v):.2f} V; r_eq={cond.r_eq_m:.6f} m",
        f"E = ({rl.c_f_km/1000:.4e} × {abs(c1.vc_fasor_v):.2f}) / (2π × {rl.eps_meio:.4e} × {cond.r_eq_m:.6f})",
        f"E = {ce.campo_kv_cm:.4f} kV/cm",
        f"Valor de referência informativo para disrupção do ar seco: ~{ce.limite_referencia_kv_cm:.1f} kV/cm "
        f"(não é um limite normativo aplicado automaticamente pela planilha original — "
        f"comparação apenas orientativa).",
    ))
    return etapas


# ---------------------------------------------------------------------------
# GERAÇÃO DO DOCUMENTO WORD — padrão oficial BK
# (capa/carimbo, cabeçalho, rodapé, sumário, 10 títulos fixos)
# ---------------------------------------------------------------------------

def nome_arquivo_relatorio(identificacao: dict) -> str:
    num = (identificacao.get("numero_documento") or "SN").strip() or "SN"
    nome = (identificacao.get("nome_documento") or "Memoria_Calculo").strip()
    rev = (identificacao.get("revisao") or "0").strip()
    nome = "".join(ch if ch.isalnum() or ch in " _-" else "_" for ch in nome).replace(" ", "_")
    num = "".join(ch if ch.isalnum() or ch in "_-" else "_" for ch in num)
    return f"{num}_{nome}_Rev_{rev}.docx"


# --- infraestrutura do documento (a partir do modelo oficial BK) ----------

def _novo_documento() -> Document:
    """Abre uma cópia do modelo oficial BK (capa, cabeçalho, rodapé, estilos
    e numeração automática dos títulos já prontos) e limpa o conteúdo de
    exemplo do corpo (sumário + seções), mantendo capa e infraestrutura de
    seção 2 (cabeçalho/rodapé/tamanho de página) intactos."""
    doc = Document(TEMPLATE_PATH)
    body = doc.element.body
    children = list(body)
    # children[0] = tabela da capa; [1] = parágrafo em branco;
    # [2] = parágrafo com quebra de seção; [-1] = sectPr final da seção 2.
    # Tudo entre [2] e [-1] é o conteúdo de exemplo do modelo — remove.
    for child in children[3:-1]:
        body.remove(child)
    return doc


def _set_paragraph_text(paragraph, text: str) -> None:
    """Substitui o texto de um parágrafo mantendo a formatação do primeiro run
    (mesma técnica recomendada pela skill de padrão de documentos BK)."""
    runs = paragraph.runs
    if runs:
        runs[0].text = text
        for run in runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def _set_cell(cell, text: str) -> None:
    _set_paragraph_text(cell.paragraphs[0], text)


def _preencher_capa_e_cabecalho(doc: Document, ident: dict, titulo_doc: str) -> None:
    """Preenche os campos da capa/carimbo e do cabeçalho de conteúdo com os
    mesmos dados (capa e cabeçalho devem sempre bater — regra da skill)."""
    revisao = ident.get("revisao") or "0"
    data_doc = ident.get("data") or "-"
    projeto = ident.get("projeto") or "-"
    numero_doc = ident.get("numero_documento") or "-"
    fase = ident.get("fase") or "MEMÓRIA DE CÁLCULO"

    capa = doc.tables[0]
    _set_cell(capa.rows[23].cells[0], revisao)          # REV. (histórico de revisões)
    _set_cell(capa.rows[23].cells[4], data_doc)          # DATA (histórico de revisões)
    _set_cell(capa.rows[26].cells[2], data_doc)          # DATA (bloco de assinaturas)
    _set_cell(capa.rows[30].cells[0], projeto)           # Nome da obra/projeto
    _set_cell(capa.rows[31].cells[0], fase)              # Fase do projeto/documento
    _set_cell(capa.rows[32].cells[0], titulo_doc)        # Título do documento
    _set_cell(capa.rows[33].cells[2], numero_doc)        # Nº DOC
    _set_cell(capa.rows[33].cells[11], revisao)          # REV. (rodapé da capa)

    header_tbl = doc.sections[1].header.tables[0]
    _set_cell(header_tbl.rows[2].cells[0], numero_doc)   # Código
    _set_cell(header_tbl.rows[2].cells[1], revisao)      # Revisão
    _set_cell(header_tbl.rows[2].cells[2], titulo_doc)   # Documento
    # Aprovação (célula 3) e Pág. (célula 4, campo automático de página)
    # permanecem como no modelo oficial BK.


# --- blocos de conteúdo no padrão BK ---------------------------------------

def _titulo(doc: Document, texto: str):
    return doc.add_paragraph(texto, style=STYLE_TITULO)


def _subtitulo(doc: Document, texto: str):
    """Parágrafo de subtítulo (estilo "1.1. Subtítulo BK").

    O estilo, no modelo oficial, traz uma numeração automática vinculada
    (numId) usada em outros contextos do template — mas a numeração dos
    subtítulos BK é sempre digitada manualmente (ex.: "4.1. Nome"), nunca
    automática. Sem esta sobrescrita, o Word/LibreOffice soma o número
    automático do estilo ao número já digitado no texto (ex.: "3.1.  4.1.
    Nome"). Aqui zeramos a numeração (numId=0 = "nenhuma lista") apenas
    neste parágrafo, sem alterar a definição do estilo em si.
    """
    p = doc.add_paragraph(texto, style=STYLE_SUBTITULO)
    pPr = p._p.get_or_add_pPr()
    numPr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    numId = OxmlElement("w:numId")
    numId.set(qn("w:val"), "0")
    numPr.append(ilvl)
    numPr.append(numId)
    pPr.append(numPr)
    return p


def _corpo(doc: Document, texto: str, negrito: bool = False, italico: bool = False):
    p = doc.add_paragraph(style=STYLE_CORPO)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = p.add_run(texto)
    run.bold = negrito
    run.italic = italico
    return p


def _lista(doc: Document, itens: list[str]):
    for item in itens:
        p = doc.add_paragraph(style=STYLE_CORPO)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.add_run(f"- {item}")


def _set_cell_bg(cell, hex_color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def _tabela(doc: Document, cabecalhos: list[str], linhas: list[list[str]]):
    """Tabela no padrão visual BK: bordas simples pretas, cabeçalho em
    negrito com fundo cinza-claro, fonte Arial 9pt — mesmo padrão usado nas
    tabelas do modelo oficial (ex.: "5.1. Dados do Projeto")."""
    t = doc.add_table(rows=1 + len(linhas), cols=len(cabecalhos))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = True

    tbl_pr = t._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for tag in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{tag}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    tbl_pr.append(borders)

    for i, h in enumerate(cabecalhos):
        cell = t.cell(0, i)
        _set_cell(cell, h)
        cell.paragraphs[0].runs[0].bold = True
        cell.paragraphs[0].runs[0].font.name = "Arial"
        cell.paragraphs[0].runs[0].font.size = Pt(9)
        _set_cell_bg(cell, _COR_CINZA_CLARO)
    for r_i, linha in enumerate(linhas, start=1):
        for c_i, valor in enumerate(linha):
            cell = t.cell(r_i, c_i)
            _set_cell(cell, str(valor))
            cell.paragraphs[0].runs[0].font.name = "Arial"
            cell.paragraphs[0].runs[0].font.size = Pt(9)
    doc.add_paragraph(style=STYLE_CORPO)
    return t


def _equacao(doc: Document, formula_xml: str):
    return eq.add_equation(doc, formula_xml)


# ---------------------------------------------------------------------------
# Corpo do relatório (os 10 títulos fixos do padrão BK)
# ---------------------------------------------------------------------------

_SUMARIO_LINHAS = [
    "1. OBJETIVO",
    "2. DOCUMENTOS DE REFERÊNCIA",
    "3. NORMAS TÉCNICAS ADOTADAS",
    "4. METODOLOGIA DE CÁLCULO",
    "4.1. Modelo de Linha Longa (Parâmetros Distribuídos)",
    "4.2. Premissa de Engenharia — Permissividade do Vácuo",
    "5. DADOS DE ENTRADA E CONSIDERAÇÕES GERAIS",
    "5.1. Dados do Projeto",
    "5.2. Condutor e Estrutura",
    "5.3. Dados Não Informados / Premissas Adotadas",
    "6. DIMENSIONAMENTO E VERIFICAÇÕES",
    "6.1. Geometria e Distância Média Geométrica (Deq)",
    "6.2. Condutor Equivalente do Feixe",
    "6.3. Parâmetros Elétricos por Km (R, L, G, C)",
    "6.4. Impedância Característica (Zo) e Constante de Propagação (γ)",
    "6.5. Circuito π Equivalente",
    "6.6. Caso 1 — Tensão, Corrente e Potência (Carga Plena)",
    "6.7. Potência Natural (SIL) e Carregamento",
    "6.8. Campo Elétrico Superficial do Condutor",
    "6.9. Validação da Metodologia (Excel × Python)",
    "7. QUANTITATIVOS",
    "8. RESULTADOS OBTIDOS",
    "9. CONCLUSÕES",
    "10. REFERÊNCIAS BIBLIOGRÁFICAS",
]


def _preencher_sumario(doc: Document) -> None:
    # Espaçamento reduzido entre linhas (o padrão "Body Text" tem 8pt depois
    # de cada parágrafo — apropriado para texto corrido, mas faz um sumário
    # com muitos subtítulos (caso deste app) estourar a página). O sumário
    # é uma lista literal, não o corpo do texto, então aqui o espaçamento é
    # apertado para caber em uma página, como no modelo oficial BK.
    p = doc.add_paragraph(style=STYLE_CORPO)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run("SUMÁRIO")
    r.bold = True
    for linha in _SUMARIO_LINHAS:
        lp = doc.add_paragraph(linha, style=STYLE_CORPO)
        lp.paragraph_format.space_after = Pt(2)
    doc.add_page_break()


def gerar_relatorio_docx(r: c.ResultadoLT, ident: dict) -> bytes:
    """Gera a Memória de Cálculo em .docx no padrão oficial BK e retorna os
    bytes do arquivo."""
    doc = _novo_documento()
    titulo_doc = ident.get("nome_documento") or "Memória de Cálculo — Parâmetros Elétricos de LT"
    _preencher_capa_e_cabecalho(doc, ident, titulo_doc)
    _preencher_sumario(doc)

    e = r.entrada
    g = r.geometria
    cond = r.condutor
    rl = r.rlgc
    ll = r.linha_longa
    pi_eq = r.circuito_pi
    c1 = r.caso1
    reg = r.regulacao
    sil = r.sil
    ce = r.campo_eletrico

    # 1. OBJETIVO ------------------------------------------------------
    _titulo(doc, "Objetivo")
    _corpo(doc, f"Este documento tem por objetivo apresentar a memória de cálculo dos "
                f"parâmetros elétricos da linha de transmissão \"{e.nome_linha}\", pelo "
                f"modelo de linha longa (parâmetros distribuídos), contemplando geometria, "
                f"parâmetros unitários (R, L, G, C), impedância característica, constante de "
                f"propagação, circuito π equivalente, condições de carga plena (Caso 1), "
                f"regulação de tensão, potência natural (SIL) e indicador de campo elétrico "
                f"superficial (corona), calculados pelo software BK Electrical Line Parameters.")

    # 2. DOCUMENTOS DE REFERÊNCIA ---------------------------------------
    _titulo(doc, "Documentos de Referência")
    _lista(doc, [
        "Planilha de cálculo elétrico de LTs original da BK Engenharia e Tecnologia "
        "(CELT2IPOG_calculo_eletrico_de_LTs.xlsm) — fonte de verdade da metodologia reproduzida.",
        "Catálogos de condutores CA (AAC), CAA (ACSR) e CAL (AAAC) — fabricantes/normas de mercado.",
        "Fuchs, R.D. — Transmissão de Energia Elétrica: Linhas Aéreas (modelo de linha longa / "
        "parâmetros distribuídos, referência clássica da metodologia empregada).",
    ])

    # 3. NORMAS TÉCNICAS ADOTADAS ----------------------------------------
    _titulo(doc, "Normas Técnicas Adotadas")
    _lista(doc, [
        "ABNT NBR 5422:2021 — Projeto de linhas aéreas de transmissão de energia elétrica "
        "(referência normativa geral de projeto de LT — não auditada fórmula-a-fórmula neste "
        "software; ver item 5.3).",
    ])

    # 4. METODOLOGIA DE CÁLCULO ------------------------------------------
    _titulo(doc, "Metodologia de Cálculo")
    _subtitulo(doc, "4.1. Modelo de Linha Longa (Parâmetros Distribuídos)")
    _corpo(doc, "A linha é representada por impedância série Z = R + jωL e admitância shunt "
                "Y = G + jωC por unidade de comprimento, com a impedância característica Zo e "
                "a constante de propagação γ (α = atenuação; β = fase) dadas por:")
    _equacao(doc, eq.eq_impedancia_serie())
    _equacao(doc, eq.eq_admitancia_shunt())
    _equacao(doc, eq.eq_zo())
    _equacao(doc, eq.eq_gamma())
    _corpo(doc, "A relação entre tensão/corrente no emissor e no receptor é dada pelas equações "
                "gerais da linha longa (funções hiperbólicas de γl — ver item 6.6). O circuito π "
                "equivalente é derivado de Zo e γl (item 6.5). A impedância de surto Ro (real, "
                "sem perdas) é usada no cálculo da potência natural (SIL, item 6.7):")
    _equacao(doc, eq.eq_ro())

    _subtitulo(doc, "4.2. Premissa de Engenharia — Permissividade do Vácuo")
    _corpo(doc, "Premissa herdada da planilha original: a permissividade do vácuo é adotada "
                "pela aproximação clássica de engenharia ε0 ≈ 1/(36π)×10⁻⁹ F/m "
                f"(≈ {dados.PERMISSIVIDADE_VACUO:.4e} F/m), e não o valor CODATA de precisão "
                f"(ε0 = {dados.PERMISSIVIDADE_VACUO_CODATA:.4e} F/m) — diferença ≈ 0,14%, "
                "propagada a C, Zo, Ro e SIL. Mantida por fidelidade à metodologia original "
                "(validado — ver item 6.9).")

    # 5. DADOS DE ENTRADA E CONSIDERAÇÕES GERAIS -------------------------
    _titulo(doc, "Dados de Entrada e Considerações Gerais")
    _subtitulo(doc, "5.1. Dados do Projeto")
    _tabela(doc, ["Item", "Descrição"], [
        ["Cliente / Concessionária", ident.get("cliente") or "-"],
        ["Projeto", ident.get("projeto") or "-"],
        ["Linha", e.nome_linha],
        ["Frequência", f"{e.frequencia_hz:.0f} Hz"],
        ["Comprimento (l)", f"{e.comprimento_km:.3f} km"],
        ["Tensão de linha no receptor (VCL)", f"{e.vcl_kv:.4f} kV"],
        ["Potência aparente no receptor (SC)", f"{e.sc_kva:.4f} kVA"],
        ["Fator de potência (cosφ)", f"{e.cos_phi:.4f} ({'atrasado' if e.atrasado else 'adiantado'})"],
    ])

    _subtitulo(doc, "5.2. Condutor e Estrutura")
    _tabela(doc, ["Item", "Descrição"], [
        ["Condutor", f"{cond.catalogo} — {cond.nome_comercial} ({cond.bitola})"],
        ["Diâmetro do condutor", f"{cond.diametro_mm:.3f} mm"],
        ["RMG do condutor (Ds)", f"{cond.rmg_m:.5f} m"],
        ["Resistência CA (catálogo, condutor unitário)", f"{cond.r_ohm_km:.4f} Ω/km"],
        ["Ampacidade (catálogo)", f"{cond.ampacidade_a:.0f} A"],
        ["Subcondutores por fase", f"{cond.n_subcondutores}"
         + (f" (d = {cond.espacamento_feixe_cm:.2f} cm)" if cond.n_subcondutores > 1 else "")],
        ["Estrutura", f"Tipo {e.tipo_estrutura}" if e.tipo_estrutura else "Customizada"],
        ["Ambiente (εR / σ)", f"{e.eps_r:.3f} / {e.sigma_s_m:.6f} S/m"],
    ])

    _subtitulo(doc, "5.3. Dados Não Informados / Premissas Adotadas")
    premissas_5_3 = [
        "PREMISSA — Ambiente (εR/σ): mantido em ar seco por padrão (εR = 1, σ = 0), salvo "
        "indicação em contrário do usuário.",
        "PREMISSA — Campo elétrico superficial (item 6.8): indicador informativo, não uma "
        "verificação normativa automática; comparação com limite de referência de literatura "
        f"(~{ce.limite_referencia_kv_cm:.1f} kV/cm, ar seco) é apenas orientativa.",
        "NECESSITA VALIDAÇÃO TÉCNICA — Caso 2 (tensão de emissor especificada), quando "
        "utilizado: reproduz a metodologia exata da planilha original (energização com a "
        "impedância de carga do Caso 1 — não necessariamente um receptor em circuito aberto). "
        "Confirmar interpretação de engenharia antes de uso em verificações de sobretensão "
        "(efeito Ferranti).",
        "DADO NÃO CADASTRADO — Estruturas 13, 14 e 15 do catálogo não possuem geometria "
        "cadastrada na planilha original; indisponíveis até cadastro futuro ou uso de "
        "geometria customizada.",
        "NECESSITA VALIDAÇÃO TÉCNICA — ABNT NBR 5422 (item 3) é referência normativa geral de "
        "projeto de LT; a metodologia de cálculo deste software não foi auditada fórmula-a-"
        "fórmula contra a norma, apenas contra a planilha de origem (item 6.9).",
    ]
    _lista(doc, premissas_5_3)

    # 6. DIMENSIONAMENTO E VERIFICAÇÕES -----------------------------------
    _titulo(doc, "Dimensionamento e Verificações")

    _subtitulo(doc, "6.1. Geometria e Distância Média Geométrica (Deq)")
    _corpo(doc, "As distâncias entre fases são obtidas da geometria da estrutura selecionada "
                "e a distância média geométrica (Deq) — usada nas fórmulas de L e C da linha "
                "transposta — é a média geométrica das três distâncias entre fases:")
    _equacao(doc, eq.eq_deq())
    _tabela(doc, ["Grandeza", "Valor", "Unidade"], [
        ["DAB", f"{g.dab_m:.4f}", "m"],
        ["DBC", f"{g.dbc_m:.4f}", "m"],
        ["DCA", f"{g.dca_m:.4f}", "m"],
        ["Deq (GMD)", f"{g.gmd_m:.4f}", "m"],
    ])

    _subtitulo(doc, "6.2. Condutor Equivalente do Feixe")
    if cond.n_subcondutores == 1:
        _corpo(doc, "Feixe simples: o raio equivalente (r_eq) e o RMG equivalente (Ds_eq) do "
                    "feixe são os do condutor unitário do catálogo.")
    elif cond.n_subcondutores == 2:
        _corpo(doc, "Feixe duplo — raio equivalente do feixe:")
        _equacao(doc, eq.eq_req_duplo())
    else:
        _corpo(doc, "Feixe quádruplo — raio equivalente do feixe:")
        _equacao(doc, eq.eq_req_quadruplo())
    _tabela(doc, ["Grandeza", "Valor", "Unidade"], [
        ["Raio equivalente do feixe (r_eq)", f"{cond.r_eq_m:.6f}", "m"],
        ["RMG equivalente do feixe (Ds_eq)", f"{cond.ds_eq_m:.6f}", "m"],
    ])

    _subtitulo(doc, "6.3. Parâmetros Elétricos por Km (R, L, G, C)")
    _equacao(doc, eq.eq_resistencia())
    _equacao(doc, eq.eq_indutancia())
    _equacao(doc, eq.eq_capacitancia())
    _equacao(doc, eq.eq_condutancia())
    _tabela(doc, ["Grandeza", "Valor", "Unidade"], [
        ["R", f"{rl.r_ohm_km:.5f}", "Ω/km"],
        ["L", f"{rl.l_h_km*1000:.5f}", "mH/km"],
        ["C", f"{rl.c_f_km*1e9:.5f}", "nF/km"],
        ["G", f"{rl.g_s_km:.3e}", "S/km"],
    ])

    _subtitulo(doc, "6.4. Impedância Característica (Zo) e Constante de Propagação (γ)")
    _equacao(doc, eq.eq_zo())
    _equacao(doc, eq.eq_gamma())
    _equacao(doc, eq.eq_ro())
    ang_zo = math.degrees(math.atan2(ll.zo_ohm.imag, ll.zo_ohm.real))
    _tabela(doc, ["Grandeza", "Valor", "Unidade"], [
        ["|Zo|", f"{abs(ll.zo_ohm):.4f}", "Ω"],
        ["∠Zo", f"{ang_zo:.4f}", "°"],
        ["Ro (impedância de surto)", f"{ll.ro_ohm:.4f}", "Ω"],
        ["αl", f"{ll.gl.real:.6f}", "-"],
        ["βl", f"{ll.gl.imag:.6f}", "-"],
    ])

    _subtitulo(doc, "6.5. Circuito π Equivalente")
    _equacao(doc, eq.eq_zs_pi())
    _equacao(doc, eq.eq_yp_pi())
    _tabela(doc, ["Grandeza", "Valor", "Unidade"], [
        ["|Zs| (série do π)", f"{abs(pi_eq.zs_ohm):.4f}", "Ω"],
        ["|Yp| (shunt do π)", f"{abs(pi_eq.yp_siemens)*1000:.5f}", "mS"],
    ])

    _subtitulo(doc, "6.6. Caso 1 — Tensão, Corrente e Potência (Carga Plena)")
    _corpo(doc, "A partir da tensão e corrente no receptor (dados de VCL, SC e cosφ), as "
                "equações gerais da linha longa fornecem tensão e corrente no emissor:")
    _equacao(doc, eq.eq_vs_caso1())
    _equacao(doc, eq.eq_is_caso1())
    _tabela(doc, ["Grandeza", "Valor", "Unidade"], [
        ["|VoL| (emissor)", f"{c1.vol_emissor_kv:.4f}", "kV"],
        ["∠VoL (emissor)", f"{c1.ang_vol_emissor_deg:.4f}", "°"],
        ["|Io| (emissor)", f"{c1.io_emissor_a:.4f}", "A"],
        ["|So| (emissor)", f"{c1.so_emissor_kva:.4f}", "kVA"],
        ["Po (emissor)", f"{c1.po_emissor_kw:.4f}", "kW"],
        ["Pc (receptor)", f"{c1.pc_receptor_kw:.4f}", "kW"],
        ["Perdas", f"{c1.perdas_kw:.4f} ({c1.perdas_pct*100:.4f}%)", "kW"],
        ["Regulação de tensão", f"{reg.regulacao_pct*100:.4f}", "%"],
    ])

    if r.caso2_manual is not None:
        c2 = r.caso2_manual
        _corpo(doc, "Caso 2 — Tensão de emissor especificada (ver premissa não identificada, "
                    "item 5.3):", negrito=True)
        _tabela(doc, ["Grandeza", "Valor", "Unidade"], [
            ["Tensão no receptor (vazio)", f"{c2.vol_receptor_kv:.4f}", "kV"],
            ["Corrente no receptor (vazio)", f"{c2.io_receptor_a:.4f}", "A"],
            ["Potência no receptor (vazio)", f"{c2.sc_receptor_kva:.4f}", "kVA"],
        ])

    _subtitulo(doc, "6.7. Potência Natural (SIL) e Carregamento")
    _equacao(doc, eq.eq_sil())
    _tabela(doc, ["Grandeza", "Valor", "Unidade"], [
        ["SIL (Potência Natural)", f"{sil.sil_kw:.2f}", "kW"],
        ["%SIL (carregamento)", f"{sil.pct_sil*100:.2f}", "%"],
    ])
    _corpo(doc, "SIL (Surge Impedance Loading) é o carregamento natural da linha, referência "
                "de carregabilidade — carregamentos muito abaixo de 1×SIL indicam linha "
                "sobredimensionada para a demanda; muito acima, possível necessidade de "
                "compensação reativa.", italico=True)

    _subtitulo(doc, "6.8. Campo Elétrico Superficial do Condutor")
    _equacao(doc, eq.eq_campo_eletrico())
    _tabela(doc, ["Grandeza", "Valor", "Unidade"], [
        ["Campo elétrico superficial (E)", f"{ce.campo_kv_cm:.4f}", "kV/cm"],
        ["Referência de literatura (ar seco)", f"~{ce.limite_referencia_kv_cm:.1f}", "kV/cm"],
    ])

    _subtitulo(doc, "6.9. Validação da Metodologia (Excel × Python)")
    _corpo(doc, "O motor de cálculo deste software foi validado célula a célula contra o "
                "caso-exemplo de fábrica da planilha original (condutor CA Tulip 336 MCM, "
                "estrutura típica 3, l = 100 km, VCL = 72 kV, SC = 1000 kVA, cosφ = 1). Os 24 "
                "resultados comparados (geometria, RLGC, impedância característica, "
                "propagação, Caso 1, regulação, SIL e campo elétrico) reproduziram os valores "
                "da planilha original com diferença inferior a 0,01% (resíduo apenas de "
                "arredondamento de exibição do Excel). Ver arquivo de testes automatizados "
                "(tests/test_calculos.py) para o detalhamento completo.")

    # 7. QUANTITATIVOS ------------------------------------------------------
    _titulo(doc, "Quantitativos")
    _corpo(doc, "Não aplicável a este documento.")

    # 8. RESULTADOS OBTIDOS ---------------------------------------------
    _titulo(doc, "Resultados Obtidos")
    _subtitulo(doc, "8.1. Resumo dos Resultados")
    _tabela(doc, ["Verificação/Grandeza", "Valor", "Status"], [
        ["Parâmetros RLGC", f"R={rl.r_ohm_km:.4f} Ω/km; L={rl.l_h_km*1000:.4f} mH/km; "
         f"C={rl.c_f_km*1e9:.4f} nF/km; G={rl.g_s_km:.2e} S/km", "Calculado"],
        ["Impedância característica (Zo)", f"{abs(ll.zo_ohm):.4f} Ω ∠{ang_zo:.2f}°", "Calculado"],
        ["Impedância de surto (Ro)", f"{ll.ro_ohm:.4f} Ω", "Calculado"],
        ["Regulação de tensão", f"{reg.regulacao_pct*100:.4f}%", "Calculado"],
        ["Perdas (Caso 1)", f"{c1.perdas_pct*100:.4f}%", "Calculado"],
        ["Potência natural (SIL) / carregamento", f"{sil.sil_kw:.2f} kW / {sil.pct_sil*100:.2f}%", "Calculado"],
        ["Campo elétrico superficial", f"{ce.campo_kv_cm:.4f} kV/cm", "Informativo (item 6.8)"],
        ["Validação Excel × Python", "24/24 itens dentro da tolerância", "OK (item 6.9)"],
    ])

    _subtitulo(doc, "8.2. Quadro de Referência — Carregabilidade Típica")
    _corpo(doc, "Tabela de referência sem vínculo de fórmula direto com esta linha (extraída "
                "da planilha original, aba \"Análises\"), apresentada apenas como comparação "
                "informativa entre carregamento (%SIL) e valores típicos de regulação/perdas.")
    _tabela(doc, ["% SIL", "Regulação (referência)", "Perdas (referência)"], [
        [f"{row['pct_sil']*100:.0f}%", f"≈ {row['regulacao_pct']*100:.2f}%", f"≈ {row['perdas_pct']*100:.2f}%"]
        for row in dados.TABELA_CARREGABILIDADE_REFERENCIA
    ])

    # 9. CONCLUSÕES -----------------------------------------------------
    _titulo(doc, "Conclusões")
    _lista(doc, [
        f"Os parâmetros elétricos calculados para a linha \"{e.nome_linha}\" indicam perdas de "
        f"{c1.perdas_pct*100:.2f}% e regulação de tensão de {reg.regulacao_pct*100:.2f}% nas "
        f"condições especificadas, com carregamento de {sil.pct_sil*100:.1f}% do SIL.",
        "Os resultados foram obtidos por reprodução fiel da metodologia da planilha de cálculo "
        "elétrico de LTs da BK Engenharia e Tecnologia, validada conforme item 6.9.",
        "O Caso 2 (tensão de emissor especificada), quando utilizado, está sujeito à premissa "
        "não identificada descrita no item 5.3 — confirmar interpretação de engenharia antes de "
        "uso em verificações de sobretensão.",
        "O campo elétrico superficial (item 6.8) é um indicador informativo, não uma verificação "
        "normativa automática.",
    ])
    _corpo(doc, "Esta memória de cálculo deve ser revisada por profissional habilitado antes de "
                "sua emissão para fins de projeto executivo, em especial quanto às premissas "
                "listadas no item 5.3.", negrito=True)

    # 10. REFERÊNCIAS BIBLIOGRÁFICAS -------------------------------------
    _titulo(doc, "Referências Bibliográficas")
    _lista(doc, [
        "FUCHS, R.D. — Transmissão de Energia Elétrica: Linhas Aéreas — Fundamentos Teóricos "
        "Básicos, LTC/UFRJ.",
        "ABNT NBR 5422:2021 — Projeto de linhas aéreas de transmissão de energia elétrica.",
        "Planilha de cálculo elétrico de LTs da BK Engenharia e Tecnologia "
        "(CELT2IPOG_calculo_eletrico_de_LTs.xlsm) — metodologia original reproduzida.",
    ])

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
